from __future__ import annotations

import os
import re
from datetime import datetime
from pathlib import Path
from threading import Lock

from flask import Flask, jsonify, render_template, send_from_directory
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
SOURCE_PATTERN = "Inspection Plan*.xlsx"
WINDOWS_SOURCE_DIR = r"C:\Friyu\RDMP RU V Balikpapan Project\QC Plan"
app = Flask(__name__, template_folder=str(BASE_DIR), static_folder=str(BASE_DIR))
_cache = {"data": None, "source": None, "mtime": None, "loaded_at": None}
_cache_lock = Lock()


def find_source() -> Path:
    configured_dir = Path(os.environ["INSPECTION_PLAN_DIR"]) if os.environ.get("INSPECTION_PLAN_DIR") else None
    mapped_windows_dir = Path("/mnt/c/Friyu/RDMP RU V Balikpapan Project/QC Plan")
    search_dirs = [directory for directory in (configured_dir, mapped_windows_dir, BASE_DIR) if directory and directory.exists()]
    candidates = [file for directory in search_dirs for file in directory.glob(SOURCE_PATTERN)]
    if not candidates:
        raise FileNotFoundError(f"Tidak ditemukan file sumber. Konfigurasi INSPECTION_PLAN_DIR untuk folder {WINDOWS_SOURCE_DIR}.")
    return candidates[0]


def text(value) -> str:
    return str(value).strip() if value is not None else ""


def work_items(detail: str) -> int:
    if not detail:
        return 0
    numbered = re.findall(r"(?:^|\n)\s*\d+\s*[.)]", detail)
    return len(numbered) or 1


def normalize_area(area: str) -> str:
    """Merge the workbook's two spellings into one reporting area."""
    normalized = re.sub(r"\s+", " ", area).strip()
    if normalized.upper() in {"PIPE SUPPORT & PLATEFORM", "PIPE SUPPORT & PLATFORM"}:
        return "PIPE SUPPORT & PLATFORM"
    return normalized


def load_plan() -> dict:
    source = find_source()
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook["Sept 2026"] if "Sept 2026" in workbook.sheetnames else workbook.worksheets[0]
    week_columns = [("W2", 4, 5), ("W3", 6, 7), ("W4", 8, 9), ("W5", 10, 11)]
    people = {}
    area_totals = {}
    job_totals = {}
    week_totals = {week: 0 for week, _, _ in week_columns}

    current_name = ""
    current_role = ""
    for row in range(6, sheet.max_row + 1):
        name = text(sheet.cell(row, 2).value)
        role = text(sheet.cell(row, 3).value)
        if name:
            current_name = name
        if role:
            current_role = role
        name = name or current_name
        role = role or current_role
        if not name or not role:
            continue
        weeks = []
        person_total = 0
        for week, area_col, detail_col in week_columns:
            area = normalize_area(text(sheet.cell(row, area_col).value))
            detail = text(sheet.cell(row, detail_col).value)
            count = work_items(detail)
            weeks.append({"week": week, "area": area, "detail": detail, "load": count})
            person_total += count
            week_totals[week] += count
            if area:
                area_entry = area_totals.setdefault(area, {"area": area, "load": 0, "people": set(), "weeks": {w: 0 for w, _, _ in week_columns}})
                area_entry["load"] += count
                area_entry["people"].add(name)
                area_entry["weeks"][week] += count
        person_key = (name, role)
        if person_key not in people:
            people[person_key] = {"name": name, "role": role, "weeks": weeks, "totalLoad": person_total}
            job_totals[role] = job_totals.get(role, 0) + person_total
        else:
            person = people[person_key]
            person["totalLoad"] += person_total
            job_totals[role] += person_total
            for existing, incoming in zip(person["weeks"], weeks):
                if incoming["area"] and existing["area"] and incoming["area"] != existing["area"]:
                    existing["area"] = f"{existing['area']} / {incoming['area']}"
                elif incoming["area"]:
                    existing["area"] = incoming["area"]
                if incoming["detail"]:
                    existing["detail"] = "\n".join(filter(None, (existing["detail"], incoming["detail"])))
                existing["load"] += incoming["load"]

    areas = []
    for item in area_totals.values():
        areas.append({"area": item["area"], "load": item["load"], "people": len(item["people"]), "weeks": item["weeks"]})
    areas.sort(key=lambda item: (-item["load"], item["area"]))
    jobs = [{"role": role, "load": load} for role, load in sorted(job_totals.items(), key=lambda pair: (-pair[1], pair[0]))]
    people = list(people.values())
    people.sort(key=lambda item: (-item["totalLoad"], item["name"]))
    source_mtime = source.stat().st_mtime
    return {
        "source": source.name,
        "sourceUpdated": datetime.fromtimestamp(source_mtime).isoformat(timespec="seconds"),
        "loadedAt": datetime.now().isoformat(timespec="seconds"),
        "weeks": [{"key": week, "label": {"W2": "W2 | 05-11 Sep", "W3": "W3 | 12-18 Sep", "W4": "W4 | 19-25 Sep", "W5": "W5 | 26 Sep-02 Oct"}[week]} for week, _, _ in week_columns],
        "people": people,
        "areas": areas,
        "jobs": jobs,
        "weekTotals": week_totals,
        "summary": {"people": len(people), "areas": len(areas), "jobs": len(jobs), "totalLoad": sum(week_totals.values())},
    }


def get_data(force: bool = False) -> dict:
    source = find_source()
    mtime = source.stat().st_mtime
    with _cache_lock:
        if force or _cache["data"] is None or _cache["mtime"] != mtime:
            _cache["data"] = load_plan()
            _cache["source"] = source.name
            _cache["mtime"] = mtime
            _cache["loaded_at"] = _cache["data"]["loadedAt"]
        return _cache["data"]


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/data")
def data():
    return jsonify(get_data())


@app.post("/api/reload")
def reload_data():
    return jsonify(get_data(force=True))


@app.get("/health")
def health():
    return {"ok": True, "source": find_source().name}


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000, debug=False)
